"""将书稿 Markdown 合成 Word 文档。

顺序：前言 → 09_开场 → 01~08章 → 10_结语 → 尾言 → 附录
"""
import re
from pathlib import Path
import md2word
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = Path(__file__).parent

# ── 文件清单（按顺序） ──────────────────────────────────
CHAPTERS_01_08 = [f"{i:02d}_" for i in range(1, 9)]
# map prefix to actual filename
PREFIX_MAP = {
    "01_": "01_互动是枢纽.md",
    "02_": "02_互动到粉丝.md",
    "03_": "03_互动与内容.md",
    "04_": "04_双路互惠.md",
    "05_": "05_消耗与缓冲.md",
    "06_": "06_崩塌与维持.md",
    "07_": "07_慢变量的力量.md",
    "08_": "08_共创与循环.md",
}

ORDER = [
    "前言.md",
    "09_开场.md",
    *[PREFIX_MAP[p] for p in CHAPTERS_01_08],
    "10_结语.md",
    "尾言.md",
    "附录_13条关系路径学科来源对照.md",
]

OUTPUT = BASE / "直播间的因果地图——六道因果络认知手册.docx"

# 在这些标题前插入分页符（正则匹配 Heading 1 文字）
PAGE_BREAK_BEFORE = [
    r"^第[一二三四五六七八九十]章",  # 第X章
    r"^第九章",
    r"^第十章",
    r"^附录",
    r"^尾言",
]


def read_md(path: Path) -> str:
    text = path.read_text(encoding="utf-8")
    return text.lstrip("\ufeff")


def concat_files(*names: str) -> str:
    parts = []
    for name in names:
        fp = BASE / name
        if not fp.exists():
            print(f"  [SKIP] 文件不存在: {name}")
            continue
        text = read_md(fp)
        # 尾言.md 如果没有 # 标题 → 补上
        if name == "尾言.md" and not text.lstrip().startswith("#"):
            text = "# 尾言\n\n" + text
        # 附录如果没有 # 标题
        if "附录" in name and not text.lstrip().startswith("#"):
            text = "# 附录：13条关系路径学科来源对照\n\n" + text
        parts.append(text)
    return "\n\n".join(parts)


def build_config() -> md2word.Config:
    config = md2word.Config()

    # 正文：宋体 小四 1.5 倍行距 首行缩进 2 字符
    config.styles["body"] = md2word.StyleConfig(
        font_name="宋体",
        font_size=12,
        first_line_indent=2,
        line_spacing_rule="multiple",
        line_spacing_value=1.5,
        space_after=6,
    )

    # 一级标题（#）：黑体 小二 加粗 居中
    config.styles["heading_1"] = md2word.StyleConfig(
        font_name="黑体",
        font_size=18,
        bold=True,
        alignment="center",
        space_before=24,
        space_after=24,
    )

    # 二级标题（##）：黑体 三号 加粗
    config.styles["heading_2"] = md2word.StyleConfig(
        font_name="黑体",
        font_size=16,
        bold=True,
        space_before=18,
        space_after=12,
    )

    # 三级标题（###）：黑体 小三 加粗
    config.styles["heading_3"] = md2word.StyleConfig(
        font_name="黑体",
        font_size=15,
        bold=True,
        space_before=12,
        space_after=6,
    )

    # 引用块：楷体 小五 左缩进
    config.styles["blockquote"] = md2word.StyleConfig(
        font_name="楷体",
        font_size=9,
        left_indent=0.5,
        first_line_indent=0,
        line_spacing_rule="multiple",
        line_spacing_value=1.3,
    )

    # 代码块：Consolas 小五
    config.styles["code"] = md2word.StyleConfig(
        font_name="Consolas",
        font_size=9,
        first_line_indent=0,
        line_spacing_rule="multiple",
        line_spacing_value=1.0,
    )

    return config


def insert_page_break_before(paragraph):
    run = paragraph.runs[0] if paragraph.runs else paragraph.add_run("")
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    run._element.insert(0, br)


def should_page_break(text: str) -> bool:
    for pat in PAGE_BREAK_BEFORE:
        if re.search(pat, text):
            return True
    return False


def post_process_docx(docx_path: Path):
    """后处理：在章节/附录/尾言标题前插入分页符。"""
    doc = Document(str(docx_path))

    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        if p.style and p.style.name == "Heading 1" and should_page_break(text):
            insert_page_break_before(p)

    doc.save(str(docx_path))


def verify_headings():
    doc = Document(str(OUTPUT))
    h1_texts = []
    for p in doc.paragraphs:
        if p.style and p.style.name == "Heading 1":
            h1_texts.append(p.text.strip())

    print(f"\n  Heading 1 数量: {len(h1_texts)}")
    for t in h1_texts:
        print(f"    [{t[:80]}]")


def main():
    print("读取 markdown 文件...")
    full_md = concat_files(*ORDER)
    print(f"  总字符数: {len(full_md):,}")

    h1_in_md = re.findall(r"^# (.+)$", full_md, re.MULTILINE)
    print(f"  # 标题数量 (md): {len(h1_in_md)}")
    for t in h1_in_md:
        print(f"    {t[:60]}")

    print("\n构建样式配置...")
    config = build_config()

    print("生成 docx（含目录）...")
    result = md2word.convert(
        full_md,
        str(OUTPUT),
        config=config,
        toc=True,
        toc_title="目录",
        toc_max_level=3,
    )
    size_kb = result.stat().st_size / 1024
    print(f"  [OK] {result.name} ({size_kb:.0f} KB)")

    print("\n后处理：插入章节分页符...")
    post_process_docx(OUTPUT)
    print("  [OK]")

    verify_headings()

    print(f"\n[DONE] 输出: {OUTPUT}")


if __name__ == "__main__":
    main()
